"""Nakliye satış KDV'sinin tevkifat sonrası raporlama testleri."""
import uuid
from datetime import date
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace

from openpyxl import load_workbook

from app.auth.models import User
from app.auth.session_security import (
    SESSION_LAST_PING_KEY,
    SESSION_TOKEN_KEY,
    new_session_token,
    utc_now,
)
from app.extensions import db
from app.firmalar.models import Firma
from app.nakliyeler.models import Nakliye
from app.services.nakliye_services import nakliye_satis_kdv_bilgisi


def _login_user(client, user_id):
    token = new_session_token()
    now = utc_now()
    user = db.session.get(User, user_id)
    user.active_session_token = token
    user.active_session_started_at = now
    user.active_session_seen_at = now
    db.session.commit()

    with client.session_transaction() as session:
        session["_user_id"] = str(user_id)
        session["_fresh"] = True
        session[SESSION_TOKEN_KEY] = token
        session[SESSION_LAST_PING_KEY] = now.isoformat()


def _firma():
    firma = Firma(
        firma_adi=f"KDV Test {uuid.uuid4().hex[:8]}",
        yetkili_adi="Yetkili",
        iletisim_bilgileri="Adres",
        vergi_dairesi="Istanbul VD",
        vergi_no=f"{uuid.uuid4().int % 10**10:010d}",
        is_musteri=True,
        is_tedarikci=True,
        bakiye=Decimal("0"),
        is_active=True,
    )
    db.session.add(firma)
    db.session.flush()
    return firma


def test_nakliye_satis_kdv_bilgisi_tevkifatli_tevkifatsiz_ve_gecersiz():
    tevkifatli = SimpleNamespace(
        kiralama_id=None,
        kiralama=None,
        kdv_orani=20,
        tevkifat_orani="2/10",
        tutar=Decimal("15000"),
    )
    bilgi = nakliye_satis_kdv_bilgisi(tevkifatli)
    assert bilgi["brut_kdv_orani"] == 20
    assert bilgi["efektif_kdv_orani"] == 16
    assert bilgi["matrah"] == 15000
    assert bilgi["kdv"] == 2400
    assert bilgi["toplam_tutar"] == 17400
    assert tevkifatli.kdv_orani == 20

    tevkifatsiz = SimpleNamespace(
        kiralama_id=None,
        kiralama=None,
        kdv_orani=20,
        tevkifat_orani=None,
        tutar=Decimal("1000"),
    )
    assert nakliye_satis_kdv_bilgisi(tevkifatsiz)["efektif_kdv_orani"] == 20

    for oran in (None, "", "2/0", "hatalı"):
        gecersiz = SimpleNamespace(
            kiralama_id=None,
            kiralama=None,
            kdv_orani=20,
            tevkifat_orani=oran,
            tutar=Decimal("1000"),
        )
        assert nakliye_satis_kdv_bilgisi(gecersiz)["efektif_kdv_orani"] == 20


def test_bagli_nakliyede_aciklamadaki_kalem_kullanilir_ilk_kalem_degil():
    ilk_kalem = SimpleNamespace(
        id=101,
        is_deleted=False,
        is_active=True,
        nakliye_satis_kdv=10,
        nakliye_satis_tevkifat_oran=None,
    )
    ikinci_kalem = SimpleNamespace(
        id=202,
        is_deleted=False,
        is_active=True,
        nakliye_satis_kdv=20,
        nakliye_satis_tevkifat_oran="2/10",
    )
    nakliye = SimpleNamespace(
        kiralama_id=1,
        kiralama=SimpleNamespace(kalemler=[ilk_kalem, ikinci_kalem]),
        aciklama="Gidiş: PF-TEST #202",
        kdv_orani=20,
        tevkifat_orani=None,
        tutar=Decimal("15000"),
    )

    bilgi = nakliye_satis_kdv_bilgisi(nakliye)
    assert bilgi["brut_kdv_orani"] == 20
    assert bilgi["tevkifat_str"] == "2/10"
    assert bilgi["efektif_kdv_orani"] == 16
    assert bilgi["kdv"] == 2400


def test_nakliye_index_yazdir_ve_excel_tevkifat_sonrasi_ayni_degeri_gosterir(app, client):
    with app.app_context():
        user = User(username=f"nak_kdv_{uuid.uuid4().hex[:8]}", rol="admin")
        user.set_password("Secret123!")
        firma = _firma()
        nakliye = Nakliye(
            firma_id=firma.id,
            tarih=date.today(),
            islem_tarihi=date.today(),
            guzergah="Yeni Magazacilik nakliyesi",
            plaka="34KDVEF16",
            tutar=Decimal("15000.00"),
            toplam_tutar=Decimal("15000.00"),
            kdv_orani=20,
            tevkifat_orani="2/10",
            is_active=True,
            is_deleted=False,
        )
        db.session.add_all([user, nakliye])
        db.session.commit()
        user_id = user.id
        firma_id = firma.id

    _login_user(client, user_id)
    query = f"?baslangic={date.today().isoformat()}&bitis={date.today().isoformat()}&firma_id={firma_id}"

    index_response = client.get(f"/nakliyeler/{query}")
    assert index_response.status_code == 200
    index_text = index_response.get_data(as_text=True)
    assert "KDV % (Tevkifat Sonrası)" in index_text
    assert "16 %" in index_text
    assert "2/10" in index_text
    assert "2.400,00 TL" in index_text
    assert "17.400,00 TL" in index_text

    print_response = client.get(f"/nakliyeler/yazdir{query}")
    assert print_response.status_code == 200
    print_text = print_response.get_data(as_text=True)
    assert "KDV % (Tevkifat Sonrası)" in print_text
    assert "%16" in print_text
    assert "2.400,00" in print_text
    assert "17.400,00" in print_text

    excel_response = client.get(f"/nakliyeler/excel{query}")
    assert excel_response.status_code == 200
    workbook = load_workbook(BytesIO(excel_response.data), read_only=True, data_only=True)
    rows = list(workbook.active.iter_rows(values_only=True))
    assert rows[3][6] == "KDV % (Tevkifat Sonrası)"
    assert rows[4][5:9] == (15000, 16, 2400, 17400)
